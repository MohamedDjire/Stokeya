"""Export Excel des statistiques."""
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework.permissions import IsAuthenticated
from rest_framework.views import APIView

from common.permissions import IsCompanyMember
from .analytics import compute_analytics

HEADER_FILL = PatternFill(start_color='1565C0', end_color='1565C0', fill_type='solid')
HEADER_FONT = Font(color='FFFFFF', bold=True)
TITLE_FONT = Font(size=14, bold=True, color='1565C0')

PERIOD_LABELS = {'week': '7 derniers jours', 'month': '30 derniers jours', 'year': 'Année en cours'}


def _style_header(ws, row, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center')


def _autosize(ws):
    for col in ws.columns:
        length = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(length + 4, 45)


class AnalyticsExportView(APIView):
    permission_classes = [IsAuthenticated, IsCompanyMember]

    def get(self, request):
        company = request.user.company
        period = request.query_params.get('period', 'month')
        data = compute_analytics(
            company,
            period=period,
            branch_id=request.query_params.get('branch') or None,
            seller_id=request.query_params.get('seller') or None,
        )

        wb = Workbook()

        # --- Feuille Synthèse ---
        ws = wb.active
        ws.title = 'Synthèse'
        ws['A1'] = f'Rapport statistique — {company.name}'
        ws['A1'].font = TITLE_FONT
        ws['A2'] = f'Période : {PERIOD_LABELS.get(data["period"], data["period"])}'
        ws['A3'] = f"Généré le {timezone.now().strftime('%d/%m/%Y %H:%M')}"

        k = data['kpis']
        ws.append([])
        ws.append(['Indicateur', 'Valeur', 'Évolution'])
        _style_header(ws, ws.max_row, 3)
        ws.append(["Chiffre d'affaires", k['revenue'], f"{k['revenue_growth']}%"])
        ws.append(['Nombre de ventes', k['sales_count'], f"{k['sales_growth']}%"])
        ws.append(['Panier moyen', k['average_basket'], ''])
        ws.append(['Produits actifs', k['active_products'], ''])
        ws.append(['Produits en stock faible', data['low_stock_count'], ''])
        _autosize(ws)

        # --- Feuille Évolution ---
        ws2 = wb.create_sheet('Évolution CA')
        ws2.append(['Période', 'Chiffre d\'affaires', 'Nb ventes'])
        _style_header(ws2, 1, 3)
        t = data['sales_trend']
        for i, label in enumerate(t['labels']):
            ws2.append([label, t['revenue'][i], t['counts'][i]])
        _autosize(ws2)

        # --- Feuille Recettes/Dépenses ---
        ws3 = wb.create_sheet('Recettes-Dépenses')
        ws3.append(['Mois', 'Recettes', 'Dépenses', 'Bénéfice'])
        _style_header(ws3, 1, 4)
        re = data['revenue_vs_expenses']
        for i, label in enumerate(re['labels']):
            ws3.append([label, re['revenue'][i], re['expenses'][i], re['profit'][i]])
        _autosize(ws3)

        # --- Feuille Top produits ---
        ws4 = wb.create_sheet('Top produits')
        ws4.append(['Produit', 'Quantité vendue', 'Chiffre d\'affaires'])
        _style_header(ws4, 1, 3)
        for p in data['top_products']:
            ws4.append([p['name'], p['quantity'], p['revenue']])
        _autosize(ws4)

        # --- Feuille Paiements & Catégories ---
        ws5 = wb.create_sheet('Répartitions')
        ws5.append(['Mode de paiement', 'Montant', 'Nb'])
        _style_header(ws5, 1, 3)
        for pm in data['payment_methods']:
            ws5.append([pm['label'], pm['total'], pm['count']])
        ws5.append([])
        cat_header = ws5.max_row + 1
        ws5.append(['Catégorie', 'Chiffre d\'affaires'])
        _style_header(ws5, cat_header, 2)
        for c in data['category_breakdown']:
            ws5.append([c['name'], c['revenue']])
        _autosize(ws5)

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        filename = f"inventra_rapport_{data['period']}_{timezone.now().strftime('%Y%m%d')}.xlsx"
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
